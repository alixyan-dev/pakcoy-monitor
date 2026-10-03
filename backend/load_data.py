#!/usr/bin/env python3
"""
Import Excel data sensor ke SQLite DB (idempoten).
Menangani 1140 baris (380 hari x 3 pengukuran), validasi, normalisasi stage.
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from app.db.database import init_db, SessionLocal, SensorReading

EXCEL = "data/new Dataset_Pakcoy.xlsx"
STAGE_MAP = {
    "Initial Stage": "initial",
    "Vegetative Stage": "vegetative",
    "Pre-Harvest Stage": "pre_harvest",
    "Harvest Ready": "harvest_ready",
}

def normalize_stage(raw: str) -> str:
    for k, v in STAGE_MAP.items():
        if k in raw:
            return v
    return "unknown"

def load():
    init_db()
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    # Cari indeks kolom berdasarkan header (case-insensitive partial)
    def idx(name):
        for i, h in enumerate(headers):
            if h and name.lower() in str(h).lower():
                return i
        return None

    i_day = idx("Day")
    i_dap = idx("DAP")
    i_time = idx("Time")
    i_moist = idx("Soil Moisture")
    i_temp = idx("Temperature")
    i_cond = idx("Soil Condition")
    i_mat = idx("Ground Truth Maturity Level")
    i_crit = idx("Ground Truth Criteria")

    if None in (i_day, i_dap, i_time, i_moist, i_temp, i_cond, i_mat, i_crit):
        raise ValueError(f"Kolom tidak lengkap. Header ditemukan: {headers}")

    db = SessionLocal()
    # Idempoten: hapus semua lalu insert ulang (untuk simplicity dalam v1)
    db.query(SensorReading).delete()
    db.commit()

    inserted = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[i_day] is None:
            continue
        db.add(SensorReading(
            day=int(row[i_day]),
            dap=int(row[i_dap]) if row[i_dap] else int(row[i_day]) + 9,
            time=str(row[i_time]).strip(),
            soil_moisture=float(row[i_moist]),
            temperature=float(row[i_temp]),
            soil_condition=str(row[i_cond]).strip(),
            maturity_pct=float(row[i_mat]),
            stage=normalize_stage(str(row[i_crit])),
        ))
        inserted += 1

    db.commit()
    db.close()
    print(f"T1.2 OK — {inserted} baris diimport ({inserted//3} hari x 3 pengukuran)")

if __name__ == "__main__":
    load()
